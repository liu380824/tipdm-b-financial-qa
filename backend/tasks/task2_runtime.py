from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Callable

import matplotlib
from openpyxl import Workbook, load_workbook

matplotlib.use('Agg')
import matplotlib.pyplot as plt

plt.rcParams['font.sans-serif'] = ['Microsoft YaHei', 'SimHei', 'DejaVu Sans']
plt.rcParams['axes.unicode_minus'] = False


QuestionRecord = dict[str, Any]
ReplyBuilder = Callable[..., dict[str, Any]]


def _normalize_batch_text(text: str) -> str:
    """Strip markdown formatting and collapse whitespace for batch export."""
    text = text.replace('{{CHART}}', '').replace('{{TABLE}}', '')
    text = re.sub(r'\{\{TASK_\d+\}\}', '', text)
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\[\u53c2\u8003\d+\]', '', text)
    text = re.sub(r'^\s*\d+\.\s*', '', text, flags=re.MULTILINE)
    text = re.sub(r'\n+', ' ', text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text


def _unique_keep_order(values: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        if value not in seen:
            seen.add(value)
            result.append(value)
    return result


def find_task2_question_file(base_dir: Path) -> Path:
    matches = [
        path
        for path in base_dir.rglob('*.xlsx')
        if '附件4' in path.name and '问题汇总' in path.name
    ]
    if not matches:
        raise FileNotFoundError('未找到任务二问题文件（附件4：问题汇总.xlsx）')
    return sorted(matches, key=lambda item: (len(item.parts), str(item)))[0]


def load_task2_questions(base_dir: Path, question_file: Path | None = None) -> list[QuestionRecord]:
    workbook_path = question_file if question_file is not None else find_task2_question_file(base_dir)
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    sheet = workbook.active
    records: list[QuestionRecord] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        question_id = str(row[0] or '').strip()
        question_type = str(row[1] or '').strip()
        raw_question = str(row[2] or '').strip()
        if not question_id or not raw_question:
            continue
        turns = json.loads(raw_question)
        normalized_turns = [
            {'Q': str(item.get('Q', '')).strip()}
            for item in turns
            if isinstance(item, dict) and str(item.get('Q', '')).strip()
        ]
        records.append({
            'id': question_id,
            'questionType': question_type,
            'questions': normalized_turns,
            'raw': raw_question
        })
    workbook.close()
    return records


_COLORS = ['#6366f1', '#0ea5e9', '#10b981', '#f59e0b', '#ef4444', '#8b5cf6', '#06b6d4', '#14b8a6']


def _extract_axis_data(option: dict[str, Any]) -> tuple[list[str], list[list[float]], list[str]]:
    """Return (labels, list_of_value_series, series_names) from ECharts option."""
    x_axis = option.get('xAxis', {})
    if isinstance(x_axis, list):
        x_axis = x_axis[0] if x_axis else {}
    labels = [str(item) for item in x_axis.get('data', [])]
    # 横向柱状图: categories are on yAxis
    if not labels:
        y_axis = option.get('yAxis', {})
        if isinstance(y_axis, list):
            y_axis = y_axis[0] if y_axis else {}
        labels = [str(item) for item in y_axis.get('data', [])]
    all_series = option.get('series', [])
    value_lists: list[list[float]] = []
    names: list[str] = []
    for s in (all_series if isinstance(all_series, list) else []):
        if not isinstance(s, dict):
            continue
        raw = s.get('data', [])
        nums = []
        for v in raw:
            if isinstance(v, (int, float)):
                nums.append(float(v))
            elif isinstance(v, str):
                nums.append(0.0)
            elif v is None:
                nums.append(0.0)
            else:
                nums.append(0.0)
        value_lists.append(nums)
        names.append(str(s.get('name', '')))
    return labels, value_lists, names


def _safe_float_list(raw: list) -> list[float]:
    out = []
    for v in raw:
        if isinstance(v, (int, float)):
            out.append(float(v))
        elif v is None or v == '-':
            out.append(0.0)
        else:
            try:
                out.append(float(v))
            except (ValueError, TypeError):
                out.append(0.0)
    return out


_PCT_KEYWORDS = {'增长率', '同比', '环比', '毛利率', '净利率', 'roe', '占比', '比率', '负债率', '%',
                 '百分比', '收益率', '利润率', '比值'}
_TURNOVER_KEYWORDS = {'周转率', '周转次数'}


def _is_pct_series(*names: str) -> bool:
    """Detect if any series/axis name implies percentage values."""
    text = ' '.join(n.lower() for n in names if n)
    return any(kw in text for kw in _PCT_KEYWORDS)


def _is_turnover_series(*names: str) -> bool:
    """Detect if any series/axis name implies turnover rate (次)."""
    text = ' '.join(n.lower() for n in names if n)
    return any(kw in text for kw in _TURNOVER_KEYWORDS)


def _fmt_val(v: float, is_pct: bool = False, is_turnover: bool = False) -> str:
    """Format a value for display on chart labels.
    If is_pct=True, format as percentage.
    If is_turnover=True, format with '次' suffix.
    Otherwise large values (>=100) are assumed to be financial amounts in 万元.
    """
    if is_turnover:
        return f'{v:.2f}次'
    if is_pct:
        av = abs(v)
        if av >= 100:
            return f'{v:,.1f}%'
        if av >= 1:
            return f'{v:.2f}%'
        return f'{v:.2f}%'
    av = abs(v)
    if av >= 1e4:
        return f'{v / 1e4:.2f}亿'
    if av >= 100:
        return f'{v:,.0f}万'
    if av >= 1:
        return f'{v:.2f}'
    if av >= 0.01:
        return f'{v:.2f}'
    if av > 0:
        return f'{v:.4f}'
    return '0'


def _gradient_colors(base_color: str, n: int) -> list[str]:
    """Generate n gradient shades of a base hex color."""
    if n <= 1:
        return [base_color]
    from matplotlib.colors import to_rgba, to_hex
    rgba = to_rgba(base_color)
    result = []
    for i in range(n):
        ratio = 0.55 + 0.45 * (1 - i / max(n - 1, 1))
        blended = tuple(c * ratio + 1.0 * (1 - ratio) for c in rgba[:3]) + (1.0,)
        result.append(to_hex(blended))
    return result


def save_chart_image(chart: dict[str, Any], output_path: Path) -> str:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    type_label = str(chart.get('typeLabel', '图表'))
    option = chart.get('option', {}) if isinstance(chart.get('option', {}), dict) else {}
    title = str(chart.get('title', output_path.stem))

    # --- professional figure style ---
    fig, ax = plt.subplots(figsize=(11, 6.2), dpi=180)
    fig.patch.set_facecolor('#fafbfc')
    ax.set_facecolor('#fafbfc')
    for spine in ax.spines.values():
        spine.set_visible(False)
    ax.tick_params(left=False, bottom=False)

    try:
        drawn = False
        # --- pie / donut ---
        if type_label in ('饼图', '环形图'):
            series = option.get('series', [])
            first = series[0] if series and isinstance(series[0], dict) else {}
            data = first.get('data', [])
            plabels = [str(d.get('name', '')) for d in data if isinstance(d, dict)]
            pvalues = [abs(float(d.get('value', 0))) for d in data if isinstance(d, dict)]
            if plabels and pvalues and len(plabels) == len(pvalues):
                wedge_kw = {'width': 0.45, 'edgecolor': 'white', 'linewidth': 2} if type_label == '环形图' else {'edgecolor': 'white', 'linewidth': 2}
                colors = _COLORS[:len(pvalues)] if len(pvalues) <= len(_COLORS) else _COLORS * (len(pvalues) // len(_COLORS) + 1)
                wedges, texts, autotexts = ax.pie(pvalues, labels=plabels, autopct='%1.1f%%',
                    colors=colors[:len(pvalues)], wedgeprops=wedge_kw, pctdistance=0.75,
                    textprops={'fontsize': 10})
                for at in autotexts:
                    at.set_fontsize(9)
                    at.set_fontweight('bold')
                drawn = True
        # --- radar ---
        elif type_label == '雷达图':
            radar_cfg = option.get('radar', {})
            indicators = radar_cfg.get('indicator', [])
            rlabels = [str(ind.get('name', '')) for ind in indicators]
            series = option.get('series', [])
            first = series[0] if series and isinstance(series[0], dict) else {}
            data_items = first.get('data', [])
            if rlabels and data_items:
                import numpy as np
                angles = np.linspace(0, 2 * np.pi, len(rlabels), endpoint=False).tolist()
                angles += angles[:1]
                ax.remove()
                ax = fig.add_subplot(111, polar=True)
                ax.set_facecolor('#fafbfc')
                ax.set_thetagrids([a * 180 / np.pi for a in angles[:-1]], rlabels, fontsize=10)
                for idx, item in enumerate(data_items[:6]):
                    vals = _safe_float_list(item.get('value', []))
                    vals += vals[:1]
                    c = _COLORS[idx % len(_COLORS)]
                    ax.plot(angles, vals, 'o-', linewidth=2.2, color=c, label=str(item.get('name', '')), markersize=5)
                    ax.fill(angles, vals, alpha=0.15, color=c)
                if len(data_items) > 1:
                    ax.legend(loc='upper right', fontsize=9, framealpha=0.9)
                drawn = True
        # --- boxplot (箱线图) ---
        elif type_label == '箱线图':
            series = option.get('series', [])
            _box_xax = option.get('xAxis', {})
            if isinstance(_box_xax, list):
                _box_xax = _box_xax[0] if _box_xax else {}
            _box_labels = _box_xax.get('data', []) if isinstance(_box_xax, dict) else []
            _box_data = []
            for s in series:
                if isinstance(s, dict) and s.get('type') == 'boxplot':
                    _box_data = s.get('data', [])
                    break
            if _box_data and _box_labels:
                # data format: each item is [min, Q1, median, Q3, max]
                bp = ax.boxplot(
                    _box_data, labels=[str(l) for l in _box_labels],
                    patch_artist=True, widths=0.4,
                    boxprops=dict(facecolor='#c7d2fe', edgecolor='#6366f1', linewidth=1.5),
                    medianprops=dict(color='#dc2626', linewidth=2),
                    whiskerprops=dict(color='#6366f1', linewidth=1.2),
                    capprops=dict(color='#6366f1', linewidth=1.2),
                    flierprops=dict(marker='o', markerfacecolor='#6366f1', markersize=5, alpha=0.5),
                )
                # color each box differently
                _box_colors = ['#c7d2fe', '#bbf7d0']
                for i, patch in enumerate(bp['boxes']):
                    patch.set_facecolor(_box_colors[i % len(_box_colors)])
                _yax_box = option.get('yAxis', {})
                if isinstance(_yax_box, list):
                    _yax_box = _yax_box[0] if _yax_box else {}
                _y_name_box = str(_yax_box.get('name', '')) if isinstance(_yax_box, dict) else ''
                if _y_name_box:
                    ax.set_ylabel(_y_name_box, fontsize=11)
                ax.grid(axis='y', linestyle='--', alpha=0.2)
                # annotate median values
                for i, d in enumerate(_box_data):
                    if isinstance(d, (list, tuple)) and len(d) >= 5:
                        median_val = d[2]
                        ax.text(i + 1, median_val, f' {median_val:.2f}%', ha='left', va='center',
                                fontsize=9, color='#dc2626', fontweight='bold')
                drawn = True
        # --- scatter / bubble ---
        elif type_label in ('散点图', '气泡图'):
            series = option.get('series', [])
            first = series[0] if series and isinstance(series[0], dict) else {}
            data = first.get('data', [])
            if data:
                def _safe_float(v, default=0):
                    try:
                        return float(v)
                    except (TypeError, ValueError):
                        return default
                xs_raw = [_safe_float(p[0]) if isinstance(p, list) and len(p) > 0 else 0 for p in data]
                ys_raw = [_safe_float(p[1]) if isinstance(p, list) and len(p) > 1 else 0 for p in data]
                sizes = [max(20, min(250, abs(_safe_float(p[2], 50)))) if isinstance(p, list) and len(p) > 2 else 50 for p in data]
                # extract axis labels from option
                _xax = option.get('xAxis', {})
                if isinstance(_xax, list):
                    _xax = _xax[0] if _xax else {}
                _yax = option.get('yAxis', {})
                if isinstance(_yax, list):
                    _yax = _yax[0] if _yax else {}
                _x_name = str(_xax.get('name', '')) if isinstance(_xax, dict) else ''
                _y_name = str(_yax.get('name', '')) if isinstance(_yax, dict) else ''
                # convert to 亿元 if values are large (万元 data)
                _x_max = max((abs(v) for v in xs_raw), default=0)
                _y_max = max((abs(v) for v in ys_raw), default=0)
                _x_unit_yi = _x_max >= 1e4
                _y_unit_yi = _y_max >= 1e4
                xs = [v / 1e4 for v in xs_raw] if _x_unit_yi else xs_raw
                ys = [v / 1e4 for v in ys_raw] if _y_unit_yi else ys_raw
                ax.scatter(xs, ys, s=sizes, alpha=0.75, color='#6366f1', edgecolors='white', linewidth=1.5)
                ax.grid(True, linestyle='--', alpha=0.2)
                # axis labels with unit
                if _x_name:
                    _xl = _x_name.replace('（万元）', '（亿元）') if _x_unit_yi else _x_name
                    ax.set_xlabel(_xl, fontsize=11, labelpad=6)
                if _y_name:
                    _yl = _y_name.replace('（万元）', '（亿元）') if _y_unit_yi else _y_name
                    ax.set_ylabel(_yl, fontsize=11, labelpad=6)
                drawn = True
        # --- heatmap ---
        elif type_label == '热力图':
            series = option.get('series', [])
            first = series[0] if series and isinstance(series[0], dict) else {}
            data = first.get('data', [])
            xl = [str(x) for x in (option.get('xAxis', {}) if not isinstance(option.get('xAxis'), list) else option.get('xAxis', [{}])[0]).get('data', [])]
            yl = [str(y) for y in (option.get('yAxis', {}) if not isinstance(option.get('yAxis'), list) else option.get('yAxis', [{}])[0]).get('data', [])]
            if xl and yl and data:
                import numpy as np
                mat = np.full((len(yl), len(xl)), float('nan'))
                for pt in data:
                    if isinstance(pt, list) and len(pt) >= 3:
                        xi, yi = int(pt[0]), int(pt[1])
                        if 0 <= xi < len(xl) and 0 <= yi < len(yl):
                            mat[yi][xi] = float(pt[2])
                im = ax.imshow(mat, aspect='auto', cmap='RdYlBu_r')
                ax.set_xticks(range(len(xl))); ax.set_xticklabels(xl, rotation=30, fontsize=9)
                ax.set_yticks(range(len(yl))); ax.set_yticklabels(yl, fontsize=9)
                fig.colorbar(im, ax=ax, shrink=0.8, pad=0.02)
                drawn = True
        # --- waterfall ---
        elif type_label == '瀑布图':
            labels, val_lists, names = _extract_axis_data(option)
            if len(val_lists) >= 3 and labels:
                base = _safe_float_list(val_lists[0])
                inc = _safe_float_list(val_lists[1])
                dec = _safe_float_list(val_lists[2])
                import numpy as np
                x = np.arange(len(labels))
                ax.bar(x, base, color='none')
                ax.bar(x, [v if v != 0 else 0 for v in inc], bottom=base, color='#10b981', label='增长', edgecolor='white')
                ax.bar(x, [v if v != 0 else 0 for v in dec], bottom=base, color='#ef4444', label='下降', edgecolor='white')
                ax.set_xticks(x); ax.set_xticklabels(labels, rotation=20, fontsize=9)
                ax.legend(fontsize=9, framealpha=0.9)
                ax.grid(axis='y', linestyle='--', alpha=0.2)
                drawn = True
        # --- histogram ---
        elif type_label == '直方图':
            labels, val_lists, _ = _extract_axis_data(option)
            if val_lists and labels:
                import numpy as np
                vals = val_lists[0]
                n = min(len(labels), len(vals))
                x = np.arange(n)
                colors = _gradient_colors('#6366f1', n)
                bars = ax.bar(x, vals[:n], color=colors, width=0.75, edgecolor='white', linewidth=0.8)
                ax.set_xticks(x)
                ax.set_xticklabels(labels[:n], rotation=30, fontsize=8, ha='right')
                ax.grid(axis='y', linestyle='--', alpha=0.2)
                for bar in bars:
                    h = bar.get_height()
                    if h > 0:
                        ax.text(bar.get_x() + bar.get_width() / 2, h, str(int(h)),
                                ha='center', va='bottom', fontsize=9, fontweight='bold', color='#4b5563')
                drawn = True
        # --- dual axis combo ---
        elif type_label == '双轴组合图':
            labels, val_lists, names = _extract_axis_data(option)
            if labels and len(val_lists) >= 2:
                n = min(len(labels), len(val_lists[0]), len(val_lists[1]))
                import numpy as np
                x = np.arange(n)
                _combo_title = str(option.get('title', {}).get('text', '')) if isinstance(option.get('title'), dict) else ''
                _combo_pct = _is_pct_series(names[0] if names else '', _combo_title)
                _combo_turnover = _is_turnover_series(names[0] if names else '', _combo_title)
                bars = ax.bar(x, val_lists[0][:n], color='#6366f1', width=0.5, label=names[0] if names else '', alpha=0.85, edgecolor='white')
                for bar in bars:
                    h = bar.get_height()
                    ax.text(bar.get_x() + bar.get_width() / 2, h, _fmt_val(h, is_pct=_combo_pct, is_turnover=_combo_turnover), ha='center', va='bottom', fontsize=7, color='#4b5563')
                ax.set_xticks(x); ax.set_xticklabels(labels[:n], rotation=20, fontsize=9)
                ax.legend(loc='upper left', fontsize=9, framealpha=0.9)
                ax.grid(axis='y', linestyle='--', alpha=0.2)
                ax2 = ax.twinx()
                for spine in ax2.spines.values():
                    spine.set_visible(False)
                ax2.plot(x, val_lists[1][:n], 'o-', color='#10b981', linewidth=2.5, label=names[1] if len(names) > 1 else '', markersize=6)
                ax2.legend(loc='upper right', fontsize=9, framealpha=0.9)
                drawn = True
        # --- treemap (fallback to horizontal bar) ---
        elif type_label == '矩形树图':
            series = option.get('series', [])
            first = series[0] if series and isinstance(series[0], dict) else {}
            data = first.get('data', [])
            tlabels = [str(d.get('name', '')) for d in data if isinstance(d, dict)]
            tvalues = [abs(float(d.get('value', 0))) for d in data if isinstance(d, dict)]
            if tlabels and tvalues:
                n = min(len(tlabels), len(tvalues))
                colors = _COLORS[:n] if n <= len(_COLORS) else _COLORS * (n // len(_COLORS) + 1)
                bars = ax.barh(tlabels[:n], tvalues[:n], color=colors[:n], height=0.6, edgecolor='white')
                for bar in bars:
                    w = bar.get_width()
                    ax.text(w, bar.get_y() + bar.get_height() / 2, f' {_fmt_val(w)}', ha='left', va='center', fontsize=9, color='#374151')
                ax.invert_yaxis()
                ax.grid(axis='x', linestyle='--', alpha=0.2)
                drawn = True
        # --- line-like (折线图/面积图/堆叠面积图) ---
        elif type_label in ('折线图', '面积图', '堆叠面积图'):
            labels, val_lists, names = _extract_axis_data(option)
            if labels and val_lists:
                _title_str = str(option.get('title', {}).get('text', '')) if isinstance(option.get('title'), dict) else ''
                _pct_ctx = _is_pct_series(*names, _title_str)
                _turnover_ctx_line = _is_turnover_series(*names, _title_str)
                for idx, vals in enumerate(val_lists):
                    n = min(len(labels), len(vals))
                    c = _COLORS[idx % len(_COLORS)]
                    lbl = names[idx] if idx < len(names) else ''
                    _pct = _pct_ctx or _is_pct_series(lbl)
                    _turnover_line = _turnover_ctx_line or _is_turnover_series(lbl)
                    ax.plot(labels[:n], vals[:n], marker='o', linewidth=2.5, color=c, label=lbl, markersize=5)
                    if type_label in ('面积图', '堆叠面积图'):
                        ax.fill_between(labels[:n], vals[:n], alpha=0.12, color=c)
                    # add value labels on data points
                    if n <= 15:
                        for i in range(n):
                            ax.annotate(_fmt_val(vals[i], is_pct=_pct, is_turnover=_turnover_line), (labels[i], vals[i]),
                                textcoords='offset points', xytext=(0, 8), ha='center', fontsize=7, color='#4b5563')
                if len(val_lists) > 1:
                    ax.legend(fontsize=9, framealpha=0.9, loc='best')
                ax.grid(axis='y', linestyle='--', alpha=0.2)
                if len(labels) > 8:
                    ax.tick_params(axis='x', rotation=30)
                drawn = True
        # --- bar-like (柱状图/横向柱状图/分组柱状图/堆叠柱状图) ---
        elif type_label in ('柱状图', '横向柱状图', '分组柱状图', '堆叠柱状图'):
            labels, val_lists, names = _extract_axis_data(option)
            if labels and val_lists:
                import numpy as np
                # detect if series are percentage-based or turnover-based from series/axis names
                _title_str = str(option.get('title', {}).get('text', '')) if isinstance(option.get('title'), dict) else ''
                _pct_ctx = _is_pct_series(*names, _title_str)
                _turnover_ctx = _is_turnover_series(*names, _title_str)
                n = min(len(labels), *(len(v) for v in val_lists))
                x = np.arange(n)
                num_series = len(val_lists)
                w = 0.7 / max(num_series, 1)
                for idx, vals in enumerate(val_lists):
                    c = _COLORS[idx % len(_COLORS)]
                    lbl = names[idx] if idx < len(names) else ''
                    _pct = _pct_ctx or _is_pct_series(lbl)
                    _turnover = _turnover_ctx or _is_turnover_series(lbl)
                    if type_label == '横向柱状图':
                        colors = _gradient_colors(c, n)
                        bars = ax.barh(x + idx * w, vals[:n], height=w, color=colors, label=lbl, edgecolor='white', linewidth=0.5)
                        _max_abs = max((abs(v) for v in vals[:n]), default=1)
                        for bar in bars:
                            bw = bar.get_width()
                            if abs(bw) < 1e-6:
                                continue
                            ax.text(bw + _max_abs * 0.01, bar.get_y() + bar.get_height() / 2,
                                f' {_fmt_val(bw, is_pct=_pct, is_turnover=_turnover)}', ha='left', va='center', fontsize=8, color='#374151', fontweight='bold')
                    elif type_label == '堆叠柱状图':
                        bottom = [sum(val_lists[j][i] for j in range(idx)) for i in range(n)] if idx > 0 else None
                        ax.bar(x, vals[:n], width=0.55, bottom=bottom, color=c, label=lbl, edgecolor='white', linewidth=0.5)
                    else:
                        bars = ax.bar(x + idx * w, vals[:n], width=w, color=c, label=lbl, edgecolor='white', linewidth=0.5)
                        # value labels on top
                        if n <= 20:
                            for bar in bars:
                                h = bar.get_height()
                                ax.text(bar.get_x() + bar.get_width() / 2, h, _fmt_val(h, is_pct=_pct, is_turnover=_turnover),
                                    ha='center', va='bottom', fontsize=7.5, color='#374151', fontweight='bold')
                if type_label == '横向柱状图':
                    ax.set_yticks(x + w * (num_series - 1) / 2)
                    ax.set_yticklabels(labels[:n], fontsize=10)
                    ax.invert_yaxis()
                    ax.grid(axis='x', linestyle='--', alpha=0.2)
                    ax.set_xticklabels([])  # data labels on bars are enough
                    ax.tick_params(axis='x', length=0)
                else:
                    ax.set_xticks(x + w * (num_series - 1) / 2)
                    ax.set_xticklabels(labels[:n], rotation=25 if n > 6 else 0, fontsize=9, ha='right' if n > 6 else 'center')
                    ax.grid(axis='y', linestyle='--', alpha=0.2)
                if num_series > 1:
                    ax.legend(fontsize=9, framealpha=0.9, loc='best')
                drawn = True

        # --- fallback: try extracting any plottable data ---
        if not drawn:
            labels, val_lists, names = _extract_axis_data(option)
            if labels and val_lists:
                n = min(len(labels), len(val_lists[0]))
                if n > 0:
                    colors = _gradient_colors('#0ea5e9', n)
                    ax.bar(labels[:n], val_lists[0][:n], color=colors, width=0.55, edgecolor='white')
                    ax.grid(axis='y', linestyle='--', alpha=0.2)
                    drawn = True
            if not drawn:
                ax.text(0.5, 0.5, type_label, ha='center', va='center', fontsize=16, color='#6b7280', transform=ax.transAxes)

        ax.set_title(title, fontsize=14, fontweight='bold', color='#1f2937', pad=16)
        fig.tight_layout(pad=1.5)
        fig.savefig(output_path, format='jpg', bbox_inches='tight', facecolor='#fafbfc')
    finally:
        plt.close(fig)
    return f'./result/{output_path.name}'


def _build_workbook(rows: list[dict[str, str]], workbook_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '任务二结果'
    header = ['编号', '问题', 'SQL 查询语句', '图形格式', '回答']
    sheet.append(header)
    for row in rows:
        sheet.append([
            row['编号'],
            row['问题'],
            row['SQL 查询语句'],
            row['图形格式'],
            row['回答']
        ])
    column_widths = {
        'A': 12,
        'B': 48,
        'C': 72,
        'D': 14,
        'E': 100
    }
    for key, width in column_widths.items():
        sheet.column_dimensions[key].width = width
    workbook.save(workbook_path)


def _process_one_question_t2(
    question: QuestionRecord,
    result_dir: Path,
    reply_builder: ReplyBuilder,
) -> tuple[dict[str, str], dict[str, Any]]:
    """Process a single task-2 question (may have multiple turns). Thread-safe."""
    import time as _time
    _t2_start = _time.time()
    question_id = str(question['id'])
    turns = list(question['questions'])
    history: list[dict[str, Any]] = []
    answers: list[dict[str, Any]] = []
    sql_fragments: list[str] = []
    chart_formats: list[str] = []
    image_index = 1

    for turn in turns:
        prompt = str(turn.get('Q', '')).strip()
        user_message = {
            'id': f'{question_id}-user-{len(history) + 1}',
            'role': 'user',
            'content': prompt
        }
        history_with_user = history + [user_message]
        reply = reply_builder(prompt, history_with_user, render_mode='batch_export', fast_mode=True, skip_kb=True)
        images: list[str] = []
        if reply.get('chart'):
            image_path = result_dir / f'{question_id}_{image_index}.jpg'
            images.append(save_chart_image(reply['chart'], image_path))
            image_index += 1
            chart_type = str(reply['chart'].get('typeLabel', '图表'))
            if chart_type:
                chart_formats.append(chart_type)
        raw_content = str(reply.get('content', ''))
        answer_payload: dict[str, Any] = {
            'content': _normalize_batch_text(raw_content)
        }
        if images:
            answer_payload['image'] = images
        answers.append({'Q': prompt, 'A': answer_payload})
        if reply.get('sql'):
            sql_fragments.append(str(reply['sql']))
        history = history_with_user + [reply]

    sql_text = '\n'.join(_unique_keep_order(sql_fragments)) if sql_fragments else '无'
    chart_format = '、'.join(_unique_keep_order(chart_formats)) if chart_formats else '无'
    answer_json = json.dumps(answers, ensure_ascii=False)
    question_json = json.dumps(turns, ensure_ascii=False)

    wb_row = {
        '编号': question_id,
        '问题': question_json,
        'SQL 查询语句': sql_text,
        '图形格式': chart_format,
        '回答': answer_json
    }
    export_rec = {
        'id': question_id,
        'questionType': question.get('questionType', ''),
        'question': turns,
        'sql': sql_text,
        'chartFormat': chart_format,
        'answer': answers
    }
    _elapsed = _time.time() - _t2_start
    _ans_preview = '\n'.join(a.get('A', {}).get('content', '')[:120] for a in answers)
    print(f'[Task2] {question_id} done in {_elapsed:.1f}s  chart={chart_format}\n  {_ans_preview}', flush=True)
    return wb_row, export_rec


def export_task2_results(base_dir: Path, reply_builder: ReplyBuilder, question_file: Path | None = None) -> dict[str, Any]:
    from concurrent.futures import ThreadPoolExecutor, as_completed
    from config import CONFIG

    questions = load_task2_questions(base_dir, question_file)
    result_dir = base_dir / 'result'
    result_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = base_dir / 'result_2.xlsx'

    max_workers = CONFIG.get('parallel', {}).get('max_workers', 3)

    # Process questions in parallel (turns within a question stay sequential)
    futures_map: dict[Any, int] = {}
    results_by_idx: dict[int, tuple[dict, dict]] = {}
    with ThreadPoolExecutor(max_workers=max_workers) as pool:
        for idx, question in enumerate(questions):
            fut = pool.submit(_process_one_question_t2, question, result_dir, reply_builder)
            futures_map[fut] = idx
        for fut in as_completed(futures_map):
            idx = futures_map[fut]
            try:
                results_by_idx[idx] = fut.result()
            except Exception as exc:
                qid = str(questions[idx].get('id', f'Q{idx}'))
                print(f'[Task2] ERROR processing {qid}: {exc}', file=__import__('sys').stderr)
                results_by_idx[idx] = (
                    {'编号': qid, '问题': '', 'SQL 查询语句': '无', '图形格式': '无', '回答': json.dumps([{'Q': '', 'A': {'content': f'处理失败: {exc}'}}], ensure_ascii=False)},
                    {'id': qid, 'questionType': '', 'question': [], 'sql': '无', 'chartFormat': '无', 'answer': []},
                )

    # Reassemble in original order
    workbook_rows = [results_by_idx[i][0] for i in range(len(questions))]
    export_records = [results_by_idx[i][1] for i in range(len(questions))]

    _build_workbook(workbook_rows, workbook_path)
    return {
        'questionFile': str(find_task2_question_file(base_dir)),
        'workbookPath': str(workbook_path),
        'workbookUrl': '/api/task2/result-file',
        'resultDir': str(result_dir),
        'questionCount': len(export_records),
        'records': export_records
    }
